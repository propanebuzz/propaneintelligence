"""Stage a complete weekly observation and repair the mislabeled demand column."""
import argparse
import json
from pathlib import Path
import subprocess
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from pipeline.storage import PRIVATE,atomic_json,archive_bytes
from pipeline.package_edit import merge_numeric
from pipeline.weekly_update import weekly_plan
from pipeline.workbooks import weekly_observations
from scripts.stage_daily_update import ROOT,NODE
from openpyxl import load_workbook


def stage(source,record,folder):
    folder.mkdir(parents=True,exist_ok=True)
    plan=weekly_plan(source,record)
    atomic_json(folder/'cell-plan.json',plan)
    if plan['cells']:
        subprocess.run([str(NODE),str(ROOT/'scripts/author_cells.mjs'),str(folder/'cell-plan.json'),str(folder/'authored-cells.xlsx')],check=True,capture_output=True)
        authored=(folder/'authored-cells.xlsx').read_bytes()
    else:authored=b''
    candidate=merge_numeric(source.read_bytes(),authored,plan)
    output=folder/'candidate.xlsx';output.write_bytes(candidate)
    before=load_workbook(source,read_only=True,data_only=False)
    after=load_workbook(output,read_only=True,data_only=False)
    for rows in zip(before['MasterData'],after['MasterData']):
        for old,new in zip(*rows):
            if old.coordinate not in plan['cells'] and old.coordinate not in plan['clear_cells'] and old.value!=new.value:raise ValueError('Unrelated weekly cell/formula changed')
    for address,expected in plan['cells'].items():
        if not address.startswith('A') and after['MasterData'][address].value!=expected:raise ValueError('Weekly numeric input readback differs: '+address)
    observed=next(r for r in weekly_observations(output) if r['date']==record['date'])
    if observed['product_supplied_bpd']!=record['product_supplied_bpd']:raise ValueError('Official demand readback failed')
    if after['MasterData'].cell(plan['row'],10).data_type=='f':raise ValueError('Official demand remains a formula')
    before.close();after.close()
    archive_bytes(PRIVATE/'archives/pre-write',source.read_bytes(),'.xlsx')
    return output,plan

if __name__=='__main__':
    report=json.loads((PRIVATE/'state/latest-run.json').read_text())
    source=PRIVATE/'archives/workbooks'/(report['drive_sources']['weekly_analytics']['sha256']+'.xlsx')
    record=json.loads((PRIVATE/'state/eia-sample-verification.json').read_text())['record']
    output,plan=stage(source,record,PRIVATE/'update-staging/weekly')
    print('Weekly candidate verified:',len(plan['cells']),'input cells;',len(plan['clear_cells']),'legacy residual formulas removed. Drive unchanged.')
