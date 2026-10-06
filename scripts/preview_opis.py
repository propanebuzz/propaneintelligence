"""Compare a privately archived OPIS PDF with daily history; never writes Drive."""
import argparse
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pdfplumber
from openpyxl import load_workbook
from pipeline.opis import parse_pages, VERSION
from pipeline.records import classify_upsert, normalize_opis


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    args = parser.parse_args()
    with pdfplumber.open(args.pdf) as document:
        candidate = parse_pages([page.extract_text() or '' for page in document.pages])
    workbook = load_workbook(args.workbook, data_only=False, read_only=True)
    sheet = workbook['Sheet1']
    if tuple(next(sheet.iter_rows(max_row=1, max_col=4, values_only=True))) != ('Date', 'Conway', 'TET', 'WTI'):
        raise ValueError('Daily workbook column layout changed')
    records, dates = [], set()
    for row, values in enumerate(sheet.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
        day, conway, tet, wti = values
        if all(v is None for v in values):
            continue
        if not isinstance(day, datetime):
            raise ValueError(f'Invalid date in daily workbook row {row}')
        record = normalize_opis({'date': day.date().isoformat(), 'conway': conway, 'tet': tet,
                                 'wti': wti, 'propane_unit': 'USD/gal', 'wti_unit': 'USD/bbl'})
        if record is None:
            raise ValueError(f'Empty dated row {row}')
        if record['date'] in dates:
            raise ValueError('Duplicate history date; reconcile before updates')
        dates.add(record['date']); records.append(record)
    workbook.close()
    preview = {'parser_version': VERSION, 'record': candidate,
               'action': classify_upsert(records, candidate), 'existing_records': len(records),
               'source_sha256': hashlib.sha256(args.pdf.read_bytes()).hexdigest(),
               'workbook_sha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
               'workbook_changed': False}
    # Fixed private destination, never the public repository or dashboard assets.
    private = Path.home() / 'Documents/Codex/propane-private/access-check'
    os.umask(0o077)
    private.mkdir(parents=True, exist_ok=True)
    (private / 'update-preview.json').write_text(json.dumps(preview, indent=2))
    (private / 'daily-normalized.json').write_text(json.dumps(records, indent=2))
    print(json.dumps(preview, indent=2))


if __name__ == '__main__':
    main()
