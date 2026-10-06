"""Complete EIA rows and removal of mislabeled legacy residual formulas."""
from datetime import date,datetime
import math
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .workbooks import FIELDS,weekly_observations

CORE={'inv':'inv_bbl','midwest':'midwest_bbl','gulf':'gulf_bbl','ready':'ready_bbl','build':'build_bbl','exports':'exports_bpd','production':'production_bpd'}


def weekly_plan(path,record):
    history=weekly_observations(path)
    day=date.fromisoformat(record['date'])
    if day.weekday()!=4 or day>date.today():raise ValueError('EIA observation must be a non-future Friday')
    mapped={get_column_letter(index+1):CORE.get(key,key) for index,key in FIELDS.items() if key!='other'}
    for key in mapped.values():
        value=record.get(key)
        if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or (value<0 and key!='build_bbl'):
            raise ValueError('Incomplete or invalid weekly input: '+key)
    workbook=load_workbook(path,read_only=True,data_only=False)
    sheet=workbook['MasterData']
    existing=next((r for r in history if r['date']==record['date']),None)
    if existing:row=existing['row']
    else:
        latest=history[-1]
        if (day-date.fromisoformat(latest['date'])).days!=7:raise ValueError('Weekly catch-up must supply each missing consecutive week')
        row=latest['row']+1
        if row>750:raise ValueError('Prefilled formula capacity exceeded; reviewed extension required')
        if sheet.cell(row,1).value is not None:raise ValueError('Next weekly row already occupied')
    clear=[]
    for cells in sheet.iter_rows(min_row=2,min_col=10,max_col=10):
        cell=cells[0];value=cell.value
        expected=f'=IF(OR($E{cell.row}="",$I{cell.row}=""),"",E{cell.row}-I{cell.row})'
        if value==expected:clear.append(cell.coordinate)
        elif isinstance(value,str) and value.startswith('='):raise ValueError('Unrecognized demand formula requires review')
    cells={}
    if not existing:cells['A'+str(row)]=(day-date(1899,12,30)).days
    for column,key in mapped.items():
        address=column+str(row);current=sheet[address].value
        if address in clear:cells[address]=record[key];continue
        if isinstance(current,str) and current.startswith('='):raise ValueError('Refusing to replace useful formula')
        if existing and current is not None and float(current)!=float(record[key]):raise ValueError('Revised EIA value requires review: '+key)
        if current!=record[key]:cells[address]=record[key]
    if not existing:
        for col in ('B','C','D','H','N','O','P','Q','R'):
            value=sheet[col+str(row)].value
            if not isinstance(value,str) or not value.startswith('='):raise ValueError('Prefilled weekly formula missing')
    workbook.close()
    return {'sheet_part':'xl/worksheets/sheet2.xml','cells':cells,'clear_cells':clear,'actions':[{'date':record['date'],'action':'insert' if not existing else ('repair_demand_column' if clear else 'duplicate')}],'row':row}
