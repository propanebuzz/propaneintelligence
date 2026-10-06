"""Read-only workbook normalization; source caches are not analytical authority."""
from collections import defaultdict
from datetime import datetime,date,timedelta
import math
from openpyxl import load_workbook
from .records import normalize_opis

FIELDS = {4:'inv',5:'midwest',6:'gulf',7:'other',8:'ready',9:'product_supplied_bpd',10:'build',11:'exports',12:'production',
          18:'product_supplied_4wk_bpd',19:'imports_bpd',20:'imports_4wk_bpd',
          21:'padd1_production_bpd',22:'padd2_production_bpd',23:'padd3_production_bpd',24:'padd45_production_bpd',
          25:'padd1_imports_bpd',26:'padd2_imports_bpd',27:'padd3_imports_bpd',28:'padd45_imports_bpd',
          29:'padd1_ready_bbl',30:'padd2_ready_bbl',31:'padd3_ready_bbl',32:'padd4_ready_bbl',33:'padd5_ready_bbl'}


def daily_prices(path):
    workbook=load_workbook(path,read_only=True,data_only=False)
    sheet=workbook['Sheet1']
    if tuple(next(sheet.iter_rows(max_row=1,max_col=4,values_only=True))) != ('Date','Conway','TET','WTI'):
        raise ValueError('Daily workbook layout changed')
    rows=[];seen=set()
    for number,(day,cwy,tet,wti) in enumerate(sheet.iter_rows(min_row=2,max_col=4,values_only=True),2):
        if all(v is None for v in (day,cwy,tet,wti)):continue
        if not isinstance(day,datetime):raise ValueError(f'Invalid daily date row {number}')
        r=normalize_opis({'date':day.date().isoformat(),'conway':cwy,'tet':tet,'wti':wti,'propane_unit':'USD/gal','wti_unit':'USD/bbl'})
        if r is None or r['date'] in seen:raise ValueError('Empty or duplicate daily history date')
        seen.add(r['date']);rows.append({'date':r['date'],'cwy':r['conway'],'tet':r['tet'],'wti':r['wti'],'row':number})
    workbook.close();return sorted(rows,key=lambda r:r['date'])


def weekly_observations(path):
    workbook=load_workbook(path,read_only=True,data_only=True)
    sheet=workbook['MasterData'];headers=next(sheet.values)
    if headers[0]!='Week Ending' or headers[9]!='EIA U.S. Propane/Propylene Product Supplied (bpd)' or headers[33]!='PADD 5 Ready-for-Sale (bbl)':
        raise ValueError('Weekly workbook layout changed')
    result=[];seen=set()
    for number,values in enumerate(sheet.iter_rows(min_row=2,max_col=34,values_only=True),2):
        day=values[0]
        if day is None:continue
        if not isinstance(day,datetime):raise ValueError('Invalid EIA date')
        iso=day.date().isoformat()
        if iso in seen:raise ValueError('Duplicate weekly source date')
        seen.add(iso);year=day.year if day.month>=4 else day.year-1
        first=date(year,4,1)
        first_friday=first+timedelta(days=(4-first.weekday())%7)
        week=round((day.date()-first_friday).days/7)+1
        r={'date':iso,'season':year,'week':week,'row':number,'cached_week':values[3]}
        for index,key in FIELDS.items():
            v=values[index]
            if v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v)):
                raise ValueError(f'Invalid weekly value {key} at row {number}')
            if v is not None and key not in ('build',) and v<0:raise ValueError('Negative stock or flow')
            r[key]=v/1e6 if v is not None and index in (4,5,6,7,8,10,11,12) else v
        if r['other'] is None and all(r[k] is not None for k in ('inv','midwest','gulf')):
            r['other']=r['inv']-r['midwest']-r['gulf']
        r['cached_yoy']=None
        result.append(r)
    workbook.close();result.sort(key=lambda r:r['date'])
    totals=defaultdict(float)
    for row in result:
        if row['build'] is not None:totals[row['season']]+=row['build']
        row['cumulative']=totals[row['season']]
    return result
